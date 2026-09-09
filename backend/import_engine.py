"""资产数据导入引擎：把上传的 Excel 工作簿解析并落库。

支持 4 种模板（自动识别）：
  - fixed_assets   固定资产清单（表头第 4 行，含「移交编号」）
                     → fixed_assets + devices + device_aliases
  - device_archive 设备档案明细（表头第 2 行，含「新设备编号」）
                     → device_archives + devices + device_accessories + device_aliases
  - ba_system      BA 系统设备清单（多 sheet：明细 + 问题清单）
                     → device_archives(BA) + ba_problems + device_aliases
  - rooms          机房信息汇总（含「机房编号」）→ rooms

设计依据：数据库设计.md v2 + XX设备档案_字段映射规范.md。
本模块**自包含**，镜像 seed_assets 的参照数据（7 子系统 + 7 BA 映射），不反向 import
asset_routes，避免循环依赖。

用法（供 HTTP 接口 / 命令行复用）：
    from import_engine import import_workbook, detect_template, TEMPLATE_INFO
    wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    result = import_workbook(db, wb, imported_by="admin", filename="xx.xlsx")
"""
from typing import Optional, Dict, Any, List
from datetime import date

import openpyxl
from sqlalchemy.orm import Session

from database import (
    Device, FixedAsset, DeviceArchive, DeviceAccessory, DeviceAlias,
    ImportBatch, EquipmentCategory, Subsystem, BaSystemMap, Room, BaProblem,
)


# ========================= 常量 =========================
# BA 明细 sheet 名（命中任一即判定为 ba_system 模板）
BA_DETAIL_SHEETS = [
    "VRV空调", "一体化空调", "一氧化碳检测", "市政排风",
    "潜污泵", "排风机", "管廊气体监测",
]
BA_PROBLEM_SHEET = "问题清单"

# 设备档案「系统」列(C) → 7 工程子系统 code（精确 + 关键词兜底）
SYSTEM_TO_SUBSYSTEM_EXACT = {
    "供电系统": "power", "电力系统": "power",
    "给排水系统": "water",
    "暖通系统": "hvac",
    "弱电系统": "weak",
    "消防系统": "fire",
    "照明系统": "lighting",
    "其他系统": "other",
}
SYSTEM_TO_SUBSYSTEM_KW = [
    ("供电", "power"), ("电力", "power"),
    ("给排水", "water"),
    ("暖通", "hvac"), ("制冷", "hvac"), ("空调", "hvac"),
    ("弱电", "weak"),
    ("消防", "fire"),
    ("照明", "lighting"), ("灯", "lighting"),
]

# 固定资产清单按文件名/段名推断子系统（覆盖常见标段关键词）
DERIVE_SUBSYSTEM_RULES = [
    ("电气", "power"), ("暖通", "hvac"), ("火灾自动报警", "fire"), ("消防", "fire"),
    ("智能化", "weak"), ("给排水", "water"), ("管道", "water"), ("给", "water"),
    ("照明", "lighting"), ("灯", "lighting"),
    ("LED屏", "other"), ("LED", "other"), ("标识", "other"), ("艺术品", "other"),
    ("柜台", "other"), ("推窗", "other"), ("排烟窗", "other"),
]

# 参照数据（镜像 asset_routes.seed_assets，保持幂等）
SUBSYSTEMS_DEF = {
    "power": ("电力系统", "Zap", 1),
    "water": ("给排水系统", "Droplets", 2),
    "hvac": ("暖通系统", "Fan", 3),
    "weak": ("弱电系统", "Cable", 4),
    "fire": ("消防系统", "Flame", 5),
    "lighting": ("照明系统", "Lightbulb", 6),
    "other": ("其他系统", "Circle", 7),
}
BA_SYSTEM_MAP_DEF = [
    ("ba_vrv", "VRV空调", "hvac", "空调类"),
    ("ba_integrated_ac", "一体化空调", "hvac", "空调类"),
    ("ba_co_detect", "一氧化碳检测", "weak", "气体监测传感器"),
    ("ba_muni_exhaust", "市政排风", "hvac", "通风"),
    ("ba_sub_pump", "潜污泵", "water", "水泵类"),
    ("ba_exhaust_fan", "排风机", "hvac", "通风"),
    ("ba_tunnel_gas", "管廊气体监测", "weak", "气体监测"),
]

TEMPLATE_INFO = {
    "fixed_assets": "固定资产清单（表头第4行，含「移交编号」）→ 固定资产总账",
    "device_archive": "设备档案明细（表头第2行，含「新设备编号」）→ 设备档案+配件",
    "ba_system": "BA系统设备清单（多sheet：明细+问题清单）→ BA设备+问题",
    "rooms": "机房信息汇总（含「机房编号」）→ 机房/房间",
    "unknown": "无法自动识别模板，请指定 source_type 或检查文件",
}


# ========================= 工具函数 =========================
def to_str(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float):
        return str(int(v)) if v == int(v) else str(v)
    if isinstance(v, (date,)):
        return v.strftime("%Y-%m-%d")
    return str(v).strip()


def parse_date(v):
    s = to_str(v).replace("．", ".").replace("。", ".")
    for sep in [".", "/", "-"]:
        if sep in s:
            parts = s.split(sep)
            if len(parts) == 3:
                try:
                    y, m, d = (int(x) for x in parts)
                    if y < 100:
                        y += 2000
                    return date(y, m, d)
                except Exception:
                    pass
    return None


def parse_float(v):
    s = to_str(v).replace(",", "").replace("，", "").replace(" ", "")
    if not s:
        return None
    try:
        return float(s)
    except Exception:
        return None


def parse_int(v):
    f = parse_float(v)
    if f is None:
        return None
    return int(f)


def _commit(db: Session, dry_run: bool):
    """dry_run 时只 flush（保留 id 供后续查询），由顶层 rollback 统一回退；否则提交。"""
    if dry_run:
        db.flush()
    else:
        db.commit()


def ensure_reference_data(db: Session, dry_run: bool = False):
    """幂等确保 7 子系统 + 7 BA 映射存在（供导入时 subsystem_id 回填）。"""
    for code, (name, icon, so) in SUBSYSTEMS_DEF.items():
        s = db.query(Subsystem).filter(Subsystem.code == code).first()
        if not s:
            s = Subsystem(code=code, name=name, icon=icon, sort_order=so)
            db.add(s)
    for code, name, sub_code, remark in BA_SYSTEM_MAP_DEF:
        m = db.query(BaSystemMap).filter(BaSystemMap.code == code).first()
        if not m:
            m = BaSystemMap(code=code, ba_system=name, subsystem_code=sub_code, remark=remark)
            db.add(m)
    _commit(db, dry_run)


def derive_subsystem_from_name(name: str) -> str:
    n = name or ""
    for kw, code in DERIVE_SUBSYSTEM_RULES:
        if kw in n:
            return code
    return "other"


def map_system_to_subsystem(system_text: str) -> str:
    s = (system_text or "").strip()
    if s in SYSTEM_TO_SUBSYSTEM_EXACT:
        return SYSTEM_TO_SUBSYSTEM_EXACT[s]
    for kw, code in SYSTEM_TO_SUBSYSTEM_KW:
        if kw in s:
            return code
    return "other"


def add_alias(db: Session, canonical: str, alias: str, source: str):
    if not alias or alias == canonical:
        return
    if db.query(DeviceAlias).filter(DeviceAlias.alias_code == alias).first():
        return
    db.add(DeviceAlias(canonical_code=canonical, alias_code=alias, source=source))


def get_or_create_batch(db: Session, fname: str, segment: str, imported_by: str, dry_run: bool = False) -> ImportBatch:
    b = db.query(ImportBatch).filter(
        ImportBatch.file_name == fname, ImportBatch.segment_name == segment,
        ImportBatch.imported_by == imported_by).first()
    if b:
        return b
    b = ImportBatch(file_name=fname, segment_name=segment, imported_by=imported_by)
    db.add(b)
    _commit(db, dry_run)
    db.refresh(b)
    return b


def upsert_category(db: Session, code: str, subsystem_code: str, dry_run: bool = False):
    if not code:
        return None
    cat = db.query(EquipmentCategory).filter(EquipmentCategory.code == code).first()
    if not cat:
        cat = EquipmentCategory(code=code, name="", subsystem_code=subsystem_code, is_active=True)
        db.add(cat)
        _commit(db, dry_run)
        db.refresh(cat)
    return cat


# ========================= 模板识别 =========================
def detect_template(wb: "openpyxl.Workbook") -> str:
    """根据 sheet 名 + 活动表表头自动识别模板类型。"""
    names = wb.sheetnames
    if "机房信息汇总" in names:
        return "rooms"
    if any(s in names for s in BA_DETAIL_SHEETS) or BA_PROBLEM_SHEET in names:
        return "ba_system"
    ws = wb.active
    for r in range(1, 9):
        joined = " ".join(str(c) for c in [ws.cell(row=r, column=c).value for c in range(1, 30)] if c)
        if "新设备编号" in joined or "原设备编号" in joined:
            return "device_archive"
        if "移交编号" in joined:
            return "fixed_assets"
    return "unknown"


# ========================= 固定资产清单导入 =========================
def import_fixed_assets_wb(db: Session, wb: "openpyxl.Workbook", imported_by: str,
                            filename: str = "", dry_run: bool = False) -> Dict[str, Any]:
    warnings: List[str] = []
    errors: List[str] = []
    batches: List[Dict[str, Any]] = []
    total = 0

    # 一个工作簿可能含多个 sheet（但通常单 sheet）；逐 sheet 处理
    for ws in wb.worksheets:
        # 自适应表头：扫描前 8 行，命中「移交编号」即表头
        hr, header, data_start = _locate_header_contains(ws, "移交编号")
        segment = ws.title or os_basename(filename)
        batch = get_or_create_batch(db, filename or ws.title, segment, imported_by, dry_run)
        batch.subsystem_hint = "fixed_assets"
        batch.area_hint = ""
        batch.status = "done"
        batch.row_count = 0
        batch.file_count = 1

        cnt = 0
        for row in ws.iter_rows(min_row=data_start, max_col=23, values_only=True):
            if row is None or all(v is None for v in row):
                continue
            vals = list(row) + [None] * (23 - len(row))
            transfer_no = to_str(vals[1])
            asset_name = to_str(vals[6])
            brand_model = to_str(vals[8])
            if not transfer_no or transfer_no == "移交编号" or "名称" in transfer_no:
                continue
            if "（" in transfer_no or "）" in transfer_no:
                continue
            if not asset_name and not brand_model:
                continue

            device_code = transfer_no
            tag_no = to_str(vals[2])
            owner_unit = to_str(vals[3])
            use_dept = to_str(vals[4])
            location = to_str(vals[5])
            asset_code = to_str(vals[7])
            serial_no = to_str(vals[9])
            recv_date = parse_date(vals[10])
            warranty_end = parse_date(vals[11])
            price_tax = parse_float(vals[12])
            price_notax = parse_float(vals[13])
            tax = parse_float(vals[14])
            budget_item = to_str(vals[15])
            contract_no = to_str(vals[16])
            bim_tag = to_str(vals[17])
            builder = to_str(vals[18])
            responsible = to_str(vals[19])
            proj_manager = to_str(vals[20])
            warranty_contact = to_str(vals[21])
            remark = to_str(vals[22])

            subsystem_code = derive_subsystem_from_name(segment)
            sub = db.query(Subsystem).filter(Subsystem.code == subsystem_code).first()
            sub_id = sub.id if sub else None
            cat = upsert_category(db, asset_code, subsystem_code, dry_run)

            dev = db.query(Device).filter(Device.device_code == device_code).first()
            if not dev:
                dev = Device(
                    device_code=device_code, name=asset_name or device_code,
                    subsystem_id=sub_id, transfer_no=transfer_no, asset_code=asset_code,
                    category_id=cat.id if cat else None, bim_tag=bim_tag,
                    system_text="", source_batch_id=batch.id, is_active=True,
                )
                db.add(dev)
            else:
                dev.subsystem_id = sub_id
                dev.transfer_no = transfer_no
                dev.asset_code = asset_code
                dev.category_id = cat.id if cat else None
                dev.bim_tag = bim_tag
                dev.source_batch_id = batch.id
                if asset_name:
                    dev.name = asset_name

            fa = db.query(FixedAsset).filter(FixedAsset.device_code == device_code).first()
            if not fa:
                fa = FixedAsset(
                    device_code=device_code, transfer_no=transfer_no, tag_no=tag_no,
                    owner_unit=owner_unit, use_dept=use_dept, location=location,
                    asset_name=asset_name, asset_code=asset_code, brand_model=brand_model,
                    serial_no=serial_no, recv_date=recv_date, warranty_end=warranty_end,
                    price_tax=price_tax, price_notax=price_notax, tax=tax,
                    budget_item=budget_item, contract_no=contract_no, bim_tag=bim_tag,
                    builder=builder, responsible=responsible, proj_manager=proj_manager,
                    warranty_contact=warranty_contact, remark=remark,
                    source_batch_id=batch.id, is_active=True,
                )
                db.add(fa)
            else:
                for k, v in {
                    "transfer_no": transfer_no, "tag_no": tag_no, "owner_unit": owner_unit,
                    "use_dept": use_dept, "location": location, "asset_name": asset_name,
                    "asset_code": asset_code, "brand_model": brand_model, "serial_no": serial_no,
                    "recv_date": recv_date, "warranty_end": warranty_end, "price_tax": price_tax,
                    "price_notax": price_notax, "tax": tax, "budget_item": budget_item,
                    "contract_no": contract_no, "bim_tag": bim_tag, "builder": builder,
                    "responsible": responsible, "proj_manager": proj_manager,
                    "warranty_contact": warranty_contact, "remark": remark,
                    "source_batch_id": batch.id,
                }.items():
                    setattr(fa, k, v)

            add_alias(db, device_code, asset_code, "fixed_asset")
            add_alias(db, device_code, bim_tag, "fixed_bim")
            add_alias(db, device_code, tag_no, "fixed_tag")
            cnt += 1

        batch.row_count = cnt
        _commit(db, dry_run)
        total += cnt
        batches.append({"segment": segment, "rows": cnt})

    return {"template": "fixed_assets", "rows": total, "batches": batches,
            "warnings": warnings, "errors": errors}


# ========================= 设备档案导入（含配件分离） =========================
def import_device_archive_wb(db: Session, wb: "openpyxl.Workbook", imported_by: str,
                             filename: str = "", dry_run: bool = False) -> Dict[str, Any]:
    warnings: List[str] = []
    errors: List[str] = []
    batches: List[Dict[str, Any]] = []
    total_main = 0
    total_acc = 0

    ws = wb.active
    # 表头第 2 行；定位含「新设备编号」的行作为表头
    hr, header, data_start = _locate_header_contains(ws, "新设备编号")
    if hr is None:
        # 兜底第 2 行
        hr, header, data_start = 2, [ws.cell(row=2, column=c).value for c in range(1, 24)], 3

    segment = os_basename(filename) or ws.title
    batch = get_or_create_batch(db, filename or ws.title, segment, imported_by, dry_run)
    batch.subsystem_hint = "device_archive"
    batch.area_hint = ""
    batch.status = "done"
    batch.row_count = 0
    batch.file_count = 1

    last_main_code: Optional[str] = None

    def add_accessory(parent_code: str, rvals, src_batch_id):
        nonlocal total_acc
        name = to_str(rvals[11])      # L 相关配件信息
        brand = to_str(rvals[12])     # M 品牌
        spec = to_str(rvals[13])      # N 规格型号
        qty = parse_int(rvals[14]) or 1   # O 数量
        unit = to_str(rvals[15])      # P 单位
        status_name = to_str(rvals[21])  # V 状态名称
        recv = parse_date(rvals[16]) or parse_date(rvals[10])  # Q 优先，否则 K
        acc = DeviceAccessory(
            parent_device_code=parent_code, name=name, brand=brand, spec=spec,
            qty=qty, unit=unit, status_name=status_name, recv_date=recv,
            source_batch_id=src_batch_id,
        )
        db.add(acc)
        total_acc += 1

    for row in ws.iter_rows(min_row=data_start, max_col=23, values_only=True):
        if row is None or all(v is None for v in row):
            continue
        vals = list(row) + [None] * (23 - len(row))
        pre_no = to_str(vals[0])        # A 预号
        project = to_str(vals[1])       # B 项目
        system_text = to_str(vals[2])   # C 系统
        location = to_str(vals[3])      # D 设备区域
        old_name = to_str(vals[4])      # E 原设备名称
        asset_name = to_str(vals[5])    # F 新设备名称
        old_code = to_str(vals[6])      # G 原设备编号
        device_code = to_str(vals[7])   # H 新设备编号（canonical）
        # vals[8] I 规格型号 —— 脏列（=原设备编号重复），丢弃
        manufacturer = to_str(vals[9])  # J 设备厂家
        recv_date = parse_date(vals[10])  # K 设备启用日期
        acc_name = to_str(vals[11])     # L 相关配件信息
        brand = to_str(vals[12])        # M 品牌
        spec = to_str(vals[13])         # N 规格型号（真实）
        qty = parse_int(vals[14]) or 1  # O 数量
        unit = to_str(vals[15]) or "台"  # P 单位
        # vals[16] Q 设备启用日期 —— 脏列（=K 重复），丢弃
        kio = to_str(vals[17])          # R KIO
        original_value = parse_float(vals[18])  # S 原值(含税)
        residual_rate = parse_float(vals[19])   # T 残值率
        net_value = parse_float(vals[20])        # U 净值
        status_name = to_str(vals[21])  # V 状态名称
        remark = to_str(vals[22])       # W 备注

        # 跳过头行（表头残留）
        if device_code == "新设备编号":
            continue
        # 模板示例行（A 列=「例如」）：仅作为后续配件的挂靠锚点，不单独建设备
        if pre_no == "例如":
            if device_code:
                last_main_code = device_code
            continue

        if device_code:
            # ---- 主设备行 ----
            brand_model = " ".join(x for x in [brand, spec] if x).strip()
            subsystem_code = map_system_to_subsystem(system_text)
            sub = db.query(Subsystem).filter(Subsystem.code == subsystem_code).first()
            sub_id = sub.id if sub else None

            dev = db.query(Device).filter(Device.device_code == device_code).first()
            if not dev:
                dev = Device(
                    device_code=device_code, name=asset_name or device_code,
                    subsystem_id=sub_id, system_text=system_text, old_code=old_code,
                    old_name=old_name, source_batch_id=batch.id, is_active=True,
                )
                db.add(dev)
            else:
                dev.subsystem_id = sub_id
                dev.system_text = system_text
                dev.old_code = old_code
                dev.old_name = old_name
                dev.source_batch_id = batch.id
                if asset_name:
                    dev.name = asset_name

            da = db.query(DeviceArchive).filter(DeviceArchive.device_code == device_code).first()
            if not da:
                da = DeviceArchive(
                    device_code=device_code, pre_no=pre_no, project=project or "T3GTC项目",
                    system_text=system_text, subsystem_id=sub_id, location=location,
                    old_name=old_name, asset_name=asset_name, old_code=old_code,
                    manufacturer=manufacturer, brand_model=brand_model, recv_date=recv_date,
                    qty=qty, unit=unit, kio=kio, original_value=original_value,
                    residual_rate=residual_rate, net_value=net_value, status_name=status_name,
                    remark=remark, template="device_archive_v1", source_batch_id=batch.id,
                    is_active=True, is_active_del=True,
                )
                db.add(da)
            else:
                da.system_text = system_text
                da.subsystem_id = sub_id
                da.location = location
                da.old_name = old_name
                da.asset_name = asset_name
                da.old_code = old_code
                da.manufacturer = manufacturer
                da.brand_model = brand_model
                da.recv_date = recv_date
                da.qty = qty
                da.unit = unit
                da.kio = kio
                da.original_value = original_value
                da.residual_rate = residual_rate
                da.net_value = net_value
                da.status_name = status_name
                da.remark = remark
                da.source_batch_id = batch.id

            add_alias(db, device_code, old_code, "archive_old")
            add_alias(db, device_code, device_code, "archive_new")
            last_main_code = device_code
            total_main += 1

            if qty and qty > 1:
                warnings.append(f"设备 {device_code} 数量={qty}，按「一台一码」仅建 1 条主设备，数量已记录待人工核对")

            # 主设备行自身若带配件（L 列非空），也作为首配件
            if acc_name:
                add_accessory(device_code, vals, batch.id)
        else:
            # ---- 配件行（H 空，L 非空）----
            if acc_name and last_main_code:
                add_accessory(last_main_code, vals, batch.id)
            elif acc_name:
                warnings.append(f"配件「{acc_name}」缺前置主设备，已跳过")
            # 其余全空行忽略

    batch.row_count = total_main
    _commit(db, dry_run)
    batches.append({"segment": segment, "rows": total_main, "accessories": total_acc})
    return {"template": "device_archive", "rows": total_main, "accessory_rows": total_acc,
            "batches": batches, "warnings": warnings, "errors": errors}


# ========================= BA 系统导入 =========================
def import_ba_wb(db: Session, wb: "openpyxl.Workbook", imported_by: str,
                 filename: str = "", dry_run: bool = False) -> Dict[str, Any]:
    warnings: List[str] = []
    errors: List[str] = []
    batches: List[Dict[str, Any]] = []
    total_dev = 0

    fname = os_basename(filename) or "BA系统设备清单"

    def bridge_device_code(ba_device_no: str):
        if not ba_device_no:
            return None
        if db.query(Device).filter(Device.device_code == ba_device_no).first():
            return ba_device_no
        a = db.query(DeviceAlias).filter(DeviceAlias.alias_code == ba_device_no).first()
        return a.canonical_code if a else None

    # 明细 sheet
    for sheet_name in BA_DETAIL_SHEETS:
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        header = [ws.cell(row=2, column=c).value for c in range(1, ws.max_column + 1)]
        maprow = db.query(BaSystemMap).filter(BaSystemMap.ba_system == sheet_name).first()
        sub_code = maprow.subsystem_code if maprow else "other"
        sub = db.query(Subsystem).filter(Subsystem.code == sub_code).first()
        sub_id = sub.id if sub else None

        batch = get_or_create_batch(db, fname, sheet_name, imported_by, dry_run)
        batch.subsystem_hint = sub_code
        batch.area_hint = ""
        batch.status = "done"
        batch.row_count = 0
        batch.file_count = 1

        cnt = 0
        for row in ws.iter_rows(min_row=3, values_only=True):
            if not row or all(v is None for v in row):
                continue
            rec = {header[i]: (row[i] if i < len(row) else None) for i in range(len(header))}
            dev_code = to_str(rec.get("设备编号"))
            if not dev_code:
                continue
            group = to_str(rec.get("组别")) or to_str(rec.get("组别/区域"))
            room = to_str(rec.get("房间"))
            loc = group
            if room:
                loc = (loc + " " + room).strip()

            d = db.query(Device).filter(Device.device_code == dev_code).first()
            if not d:
                d = Device(device_code=dev_code, name=dev_code, subsystem_id=sub_id,
                           system_text=sheet_name, source_batch_id=batch.id)
                db.add(d)
            else:
                d.subsystem_id = sub_id
                d.system_text = sheet_name
                d.source_batch_id = batch.id

            da = db.query(DeviceArchive).filter(DeviceArchive.device_code == dev_code).first()
            if not da:
                da = DeviceArchive(
                    device_code=dev_code, system_text=sheet_name, subsystem_id=sub_id,
                    location=loc, project="T3GTC项目", template="ba_device_v1", status_name="",
                    is_active_del=True, source_batch_id=batch.id, extra=rec, is_active=True,
                    qty=1, unit="台",
                )
                db.add(da)
            else:
                da.subsystem_id = sub_id
                da.system_text = sheet_name
                da.location = loc
                da.extra = rec
                da.source_batch_id = batch.id

            add_alias(db, dev_code, dev_code, "ba_device")
            cnt += 1
        batch.row_count = cnt
        _commit(db, dry_run)
        total_dev += cnt
        batches.append({"segment": sheet_name, "rows": cnt})

    # 问题清单 sheet
    if BA_PROBLEM_SHEET in wb.sheetnames:
        ws = wb[BA_PROBLEM_SHEET]
        header = [ws.cell(row=3, column=c).value for c in range(1, ws.max_column + 1)]
        batch = get_or_create_batch(db, fname, BA_PROBLEM_SHEET, imported_by, dry_run)
        batch.subsystem_hint = "ba"
        batch.area_hint = ""
        batch.status = "done"
        batch.row_count = 0
        batch.file_count = 1

        # 幂等：先清本批次已有问题
        db.query(BaProblem).filter(BaProblem.source_batch_id == batch.id).delete()
        db.flush()

        cnt = 0
        for row in ws.iter_rows(min_row=4, values_only=True):
            if not row or all(v is None for v in row):
                continue
            rec = {header[i]: (row[i] if i < len(row) else None) for i in range(len(header))}
            ba_system = to_str(rec.get("所属系统"))
            ba_device_no = to_str(rec.get("设备编号"))
            group_area = to_str(rec.get("组别/区域"))
            location = to_str(rec.get("位置"))
            problem_type = to_str(rec.get("问题类型"))
            maprow = db.query(BaSystemMap).filter(BaSystemMap.ba_system == ba_system).first()
            ba_system_code = maprow.code if maprow else None
            device_code = bridge_device_code(ba_device_no)
            exist = db.query(BaProblem).filter(
                BaProblem.ba_device_no == ba_device_no, BaProblem.ba_system == ba_system).first()
            if exist:
                exist.device_code = device_code
                exist.ba_system_code = ba_system_code
                exist.problem_type = problem_type
                exist.group_area = group_area
                exist.location = location
                exist.source_batch_id = batch.id
            else:
                db.add(BaProblem(
                    ba_device_no=ba_device_no, device_code=device_code, ba_system=ba_system,
                    ba_system_code=ba_system_code, problem_type=problem_type,
                    group_area=group_area, location=location, status="open",
                    source_batch_id=batch.id,
                ))
                cnt += 1
        batch.row_count = db.query(BaProblem).filter(BaProblem.source_batch_id == batch.id).count()
        _commit(db, dry_run)
        batches.append({"segment": BA_PROBLEM_SHEET, "rows": batch.row_count})

    return {"template": "ba_system", "rows": total_dev, "batches": batches,
            "warnings": warnings, "errors": errors}


# ========================= 机房导入 =========================
def import_rooms_wb(db: Session, wb: "openpyxl.Workbook", imported_by: str,
                    filename: str = "", dry_run: bool = False) -> Dict[str, Any]:
    warnings: List[str] = []
    errors: List[str] = []
    batches: List[Dict[str, Any]] = []

    ws = wb["机房信息汇总"] if "机房信息汇总" in wb.sheetnames else wb.active
    # 定位含「机房编号」的表头行（前 8 行）
    hr, header, _ = _locate_header_contains(ws, "机房编号")
    if hr is None:
        hr = 3
    idx = {}
    for i, h in enumerate(header or [], 1):
        if isinstance(h, str):
            for key in ["楼栋", "楼层", "机房名称", "机房编号"]:
                if key in h:
                    idx[key] = i
    if not all(k in idx for k in ["楼栋", "楼层", "机房名称", "机房编号"]):
        return {"template": "rooms", "rows": 0, "batches": [],
                "warnings": warnings, "errors": ["未找到机房表头列（楼栋/楼层/机房名称/机房编号）"]}

    batch = get_or_create_batch(db, os_basename(filename) or ws.title, "机房信息汇总", imported_by, dry_run)
    batch.subsystem_hint = "rooms"
    batch.area_hint = ""
    batch.status = "done"
    batch.row_count = 0
    batch.file_count = 1

    cnt = 0
    for row in ws.iter_rows(min_row=hr + 1, values_only=True):
        if not row:
            continue
        def cell(key):
            c = idx[key]
            return to_str(row[c - 1]) if len(row) >= c else ""
        code = cell("机房编号")
        if not code:
            continue
        building = cell("楼栋")
        floor = cell("楼层")
        name = cell("机房名称")
        room = db.query(Room).filter(Room.code == code).first()
        if not room:
            room = Room(code=code, building=building, floor=floor, name=name,
                        room_type="高频", is_active=True, shift="morning")
            db.add(room)
        cnt += 1
    batch.row_count = cnt
    _commit(db, dry_run)
    batches.append({"segment": "机房信息汇总", "rows": cnt})
    return {"template": "rooms", "rows": cnt, "batches": batches,
            "warnings": warnings, "errors": errors}


# ========================= 内部工具 =========================
def os_basename(p: str) -> str:
    if not p:
        return ""
    return p.replace("\\", "/").split("/")[-1]


def _locate_header_contains(ws, marker: str):
    """扫描前 8 行，返回 (header_row, header_cells, data_start_row)。找不到返回 (None,None,9)。"""
    for r in range(1, 9):
        cells = [ws.cell(row=r, column=c).value for c in range(1, 24)]
        for c in cells:
            if isinstance(c, str) and marker in c:
                return r, cells, r + 1
    return None, None, 9


# ========================= 统一入口 =========================
def import_workbook(db: Session, wb: "openpyxl.Workbook", imported_by: str = "api",
                    source_type: Optional[str] = None, filename: str = "",
                    dry_run: bool = False) -> Dict[str, Any]:
    """解析并导入一个工作簿。返回结构化结果；dry_run=True 时解析并计数但不落库。"""
    prev_autoflush = db.autoflush
    db.autoflush = True
    try:
        ensure_reference_data(db, dry_run=dry_run)
        tpl = source_type or detect_template(wb)
        if tpl == "fixed_assets":
            res = import_fixed_assets_wb(db, wb, imported_by, filename, dry_run)
        elif tpl == "device_archive":
            res = import_device_archive_wb(db, wb, imported_by, filename, dry_run)
        elif tpl == "ba_system":
            res = import_ba_wb(db, wb, imported_by, filename, dry_run)
        elif tpl == "rooms":
            res = import_rooms_wb(db, wb, imported_by, filename, dry_run)
        else:
            raise ValueError(f"无法识别模板：{tpl}")
        if dry_run:
            db.rollback()
        return res
    finally:
        db.autoflush = prev_autoflush

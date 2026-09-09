"""导入机房信息汇总（516 间）到 rooms 表。

权威源：D:/机场业务/T3GTC/02_设备清单/机房信息汇总(PDF版).xlsx
  sheet「机房信息汇总」共 516 间：东停车楼190 / 西停车楼186 / GTC140。
  列（第 3 行为表头）：楼栋 / 楼层 / 机房名称 / 机房编号。

说明：本库 app.db 的 rooms 实际为空（设计文档称已种子，但本环境未落库），
故补一个幂等导入脚本，使区域树 / location→room 归一化可用。
按 机房编号(code) upsert。

用法：
  venv/Scripts/python.exe scripts/import_rooms.py
"""
import os
import sys
import argparse

import openpyxl

BACKEND_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, BACKEND_DIR)

from database import engine, Room, ImportBatch
from sqlalchemy.orm import sessionmaker

Session = sessionmaker(bind=engine, autoflush=True)

DEFAULT_FILE = r"D:/机场业务/T3GTC/02_设备清单/机房信息汇总(PDF版).xlsx"
DEFAULT_SHEET = "机房信息汇总"


def to_str(v):
    if v is None:
        return ""
    if isinstance(v, float):
        return str(int(v)) if v == int(v) else str(v)
    return str(v).strip()


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--file", default=DEFAULT_FILE)
    ap.add_argument("--sheet", default=DEFAULT_SHEET)
    ap.add_argument("--header-row", type=int, default=3)
    ap.add_argument("--force", action="store_true", help="覆盖已存在机房的名称/楼栋/楼层")
    args = ap.parse_args()

    db = Session()
    try:
        wb = openpyxl.load_workbook(args.file, read_only=True, data_only=True)
        ws = wb[args.sheet]
        # 定位列：按表头名称映射
        header = [ws.cell(row=args.header_row, column=c).value for c in range(1, 8)]
        idx = {}
        for i, h in enumerate(header, 1):
            if isinstance(h, str):
                for key in ["楼栋", "楼层", "机房名称", "机房编号"]:
                    if key in h:
                        idx[key] = i
        print("列映射:", idx)
        assert all(k in idx for k in ["楼栋", "楼层", "机房名称", "机房编号"]), "表头列缺失"

        batch = db.query(ImportBatch).filter(
            ImportBatch.segment_name == "机房信息汇总", ImportBatch.imported_by == "seed_rooms").first()
        if not batch:
            batch = ImportBatch(segment_name="机房信息汇总", imported_by="seed_rooms", status="done")
            db.add(batch); db.commit(); db.refresh(batch)

        cnt = 0
        for row in ws.iter_rows(min_row=args.header_row + 1, values_only=True):
            if not row:
                continue
            code = to_str(row[idx["机房编号"] - 1]) if len(row) >= idx["机房编号"] else ""
            if not code:
                continue
            building = to_str(row[idx["楼栋"] - 1]) if len(row) >= idx["楼栋"] else ""
            floor = to_str(row[idx["楼层"] - 1]) if len(row) >= idx["楼层"] else ""
            name = to_str(row[idx["机房名称"] - 1]) if len(row) >= idx["机房名称"] else ""
            room = db.query(Room).filter(Room.code == code).first()
            if not room:
                room = Room(code=code, building=building, floor=floor, name=name,
                            room_type="高频", is_active=True, shift="morning")
                db.add(room)
                cnt += 1
            elif args.force:
                room.building = building
                room.floor = floor
                room.name = name
        db.commit()
        total = db.query(Room).count()
        batch.row_count = total
        batch.file_count = 1
        db.commit()
        print(f"机房导入完成：本次新增 {cnt} 间，rooms 总计 {total} 间。")
    finally:
        db.close()


if __name__ == "__main__":
    main()

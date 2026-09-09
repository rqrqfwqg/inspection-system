"""导入固定资产清单（46 个 xlsx，表头第 4 行、23 列）到 fixed_assets + devices + device_aliases。

每个文件 = 1 条 ImportBatch；按 device_code（= 移交编号）幂等 upsert。
用法：
  venv/Scripts/python.exe scripts/import_fixed_assets.py            # 全量
  venv/Scripts/python.exe scripts/import_fixed_assets.py --limit 2  # 仅前 2 个文件（验证）
  venv/Scripts/python.exe scripts/import_fixed_assets.py --dir <路径>
"""
import os
import sys
import re
import argparse
from datetime import datetime, timezone, date

import openpyxl

# 让脚本能 import 后端模块
BACKEND_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, BACKEND_DIR)

from database import (
    engine, Device, FixedAsset, DeviceAlias, ImportBatch,
    EquipmentCategory, Subsystem,
)
from sqlalchemy.orm import sessionmaker

# 使用 autoflush=True 的会话：保证 upsert 时能看到本文件内已 add 但未 commit 的记录
# （database.SessionLocal 为 autoflush=False，同文件内重复 device_code 会误判为「新增」导致 UNIQUE 冲突）。
Session = sessionmaker(bind=engine, autoflush=True)

DEFAULT_DIR = r"D:/机场业务/T3GTC/03_资产清单/公管分公司资产清单/公管分公司资产清单/"


# ========================= 解析工具 =========================
def to_str(v):
    if v is None:
        return ""
    if isinstance(v, float):
        return str(int(v)) if v == int(v) else str(v)
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y-%m-%d")
    return str(v).strip()


def parse_date(v):
    s = to_str(v).replace("．", ".").replace("。", ".")
    for sep in [".", "/", "-"]:
        if sep in s:
            parts = re.split(r"[./-]", s)
            if len(parts) == 3:
                try:
                    y, m, d = (int(x) for x in parts)
                    if y < 100:
                        y += 2000
                    return date(y, m, d)
                except Exception:
                    pass
    return None


def parse_float(v):
    s = to_str(v).replace(",", "").replace("，", "").replace(" ", "")
    if not s:
        return None
    try:
        return float(s)
    except Exception:
        return None


def derive_subsystem(name: str) -> str:
    """从文件名/段名推断子系统 code（覆盖常见标段关键词）。"""
    n = name or ""
    rules = [
        ("电气", "power"),
        ("暖通", "hvac"),
        ("火灾自动报警", "fire"),
        ("消防", "fire"),
        ("智能化", "weak"),
        ("给排水", "water"),
        ("管道", "water"),
        ("给", "water"),
        ("照明", "lighting"),
        ("灯", "lighting"),
        ("LED屏", "other"),
        ("LED", "other"),
        ("标识", "other"),
        ("艺术品", "other"),
        ("柜台", "other"),
        ("推窗", "other"),
        ("排烟窗", "other"),
    ]
    for kw, code in rules:
        if kw in n:
            return code
    return "other"


def add_alias(db, canonical: str, alias: str, source: str):
    """get-or-create：alias_code 唯一，重复跳过。"""
    if not alias or alias == canonical:
        return
    exists = db.query(DeviceAlias).filter(DeviceAlias.alias_code == alias).first()
    if exists:
        return
    db.add(DeviceAlias(canonical_code=canonical, alias_code=alias, source=source))


def detect_header(ws):
    """自适应定位表头行：扫描前 8 行，命中含「移交编号」的行即表头。"""
    for r in range(1, 9):
        cells = [ws.cell(row=r, column=c).value for c in range(1, 24)]
        for c in cells:
            if isinstance(c, str) and "移交编号" in c:
                return r, cells, r + 1
    # 兜底：默认第 4 行
    return 4, [ws.cell(row=4, column=c).value for c in range(1, 24)], 5


def get_or_create_batch(db, fname, segment, imported_by):
    """按 (file_name, segment_name, imported_by) 幂等复用批次。"""
    b = db.query(ImportBatch).filter(
        ImportBatch.file_name == fname, ImportBatch.segment_name == segment,
        ImportBatch.imported_by == imported_by).first()
    if b:
        return b
    b = ImportBatch(file_name=fname, segment_name=segment, imported_by=imported_by)
    db.add(b)
    db.commit()
    db.refresh(b)
    return b


def upsert_category(db, code: str, subsystem_code: str):
    if not code:
        return None
    cat = db.query(EquipmentCategory).filter(EquipmentCategory.code == code).first()
    if not cat:
        cat = EquipmentCategory(code=code, name="", subsystem_code=subsystem_code, is_active=True)
        db.add(cat)
        db.commit()
        db.refresh(cat)
    return cat


# ========================= 单文件导入 =========================
def import_file(db, path: str):
    fname = os.path.basename(path)
    segment = os.path.splitext(fname)[0]
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.active

    batch = get_or_create_batch(db, fname, segment, "import_fixed_assets")
    batch.subsystem_hint = derive_subsystem(fname)
    batch.area_hint = ""
    batch.status = "done"
    batch.row_count = 0
    batch.file_count = 1

    # 自适应表头定位（不同文件表头行可能不同）
    hr, header, data_start = detect_header(ws)
    subsystem_code = derive_subsystem(fname)
    sub = db.query(Subsystem).filter(Subsystem.code == subsystem_code).first()
    sub_id = sub.id if sub else None

    count = 0
    for row in ws.iter_rows(min_row=data_start, max_col=23, values_only=True):
        if row is None or all(v is None for v in row):
            continue
        vals = list(row) + [None] * (23 - len(row))
        transfer_no = to_str(vals[1])
        asset_name = to_str(vals[6])
        brand_model = to_str(vals[8])
        # 跳过非数据行：空 / 表头残留（含「名称」） / 章节标记 / 无资产名且无品牌的占位行
        if not transfer_no or transfer_no == "移交编号" or "名称" in transfer_no:
            continue
        if "（" in transfer_no or "）" in transfer_no:
            continue
        if not asset_name and not brand_model:
            continue
        device_code = transfer_no  # 移交编号即 device_code（检索主键）
        tag_no = to_str(vals[2])
        owner_unit = to_str(vals[3])
        use_dept = to_str(vals[4])
        location = to_str(vals[5])
        asset_name = to_str(vals[6])
        asset_code = to_str(vals[7])
        brand_model = to_str(vals[8])
        serial_no = to_str(vals[9])
        recv_date = parse_date(vals[10])
        warranty_end = parse_date(vals[11])
        price_tax = parse_float(vals[12])
        price_notax = parse_float(vals[13])
        tax = parse_float(vals[14])
        budget_item = to_str(vals[15])
        contract_no = to_str(vals[16])
        bim_tag = to_str(vals[17])
        builder = to_str(vals[18])
        responsible = to_str(vals[19])
        proj_manager = to_str(vals[20])
        warranty_contact = to_str(vals[21])
        remark = to_str(vals[22])

        cat = upsert_category(db, asset_code, subsystem_code)

        # Device upsert
        dev = db.query(Device).filter(Device.device_code == device_code).first()
        if not dev:
            dev = Device(
                device_code=device_code, name=asset_name or device_code,
                subsystem_id=sub_id, transfer_no=transfer_no, asset_code=asset_code,
                category_id=cat.id if cat else None, bim_tag=bim_tag,
                system_text="", source_batch_id=batch.id, is_active=True,
            )
            db.add(dev)
        else:
            dev.subsystem_id = sub_id
            dev.transfer_no = transfer_no
            dev.asset_code = asset_code
            dev.category_id = cat.id if cat else None
            dev.bim_tag = bim_tag
            dev.source_batch_id = batch.id
            if asset_name:
                dev.name = asset_name

        # FixedAsset upsert
        fa = db.query(FixedAsset).filter(FixedAsset.device_code == device_code).first()
        if not fa:
            fa = FixedAsset(
                device_code=device_code, transfer_no=transfer_no, tag_no=tag_no,
                owner_unit=owner_unit, use_dept=use_dept, location=location,
                asset_name=asset_name, asset_code=asset_code, brand_model=brand_model,
                serial_no=serial_no, recv_date=recv_date, warranty_end=warranty_end,
                price_tax=price_tax, price_notax=price_notax, tax=tax,
                budget_item=budget_item, contract_no=contract_no, bim_tag=bim_tag,
                builder=builder, responsible=responsible, proj_manager=proj_manager,
                warranty_contact=warranty_contact, remark=remark,
                source_batch_id=batch.id, is_active=True,
            )
            db.add(fa)
        else:
            for k, v in {
                "transfer_no": transfer_no, "tag_no": tag_no, "owner_unit": owner_unit,
                "use_dept": use_dept, "location": location, "asset_name": asset_name,
                "asset_code": asset_code, "brand_model": brand_model, "serial_no": serial_no,
                "recv_date": recv_date, "warranty_end": warranty_end, "price_tax": price_tax,
                "price_notax": price_notax, "tax": tax, "budget_item": budget_item,
                "contract_no": contract_no, "bim_tag": bim_tag, "builder": builder,
                "responsible": responsible, "proj_manager": proj_manager,
                "warranty_contact": warranty_contact, "remark": remark,
                "source_batch_id": batch.id,
            }.items():
                setattr(fa, k, v)

        # Aliases：资产代码 / BIM标签 / 标签号
        add_alias(db, device_code, asset_code, "fixed_asset")
        add_alias(db, device_code, bim_tag, "fixed_bim")
        add_alias(db, device_code, tag_no, "fixed_tag")

        count += 1

    batch.row_count = count
    db.commit()
    wb.close()
    return count


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--dir", default=DEFAULT_DIR)
    ap.add_argument("--limit", type=int, default=0, help="仅导入前 N 个文件（验证用）")
    args = ap.parse_args()

    files = sorted([os.path.join(args.dir, f) for f in os.listdir(args.dir)
                    if f.lower().endswith(".xlsx")])
    if args.limit:
        files = files[:args.limit]

    db = Session()
    total = 0
    try:
        for i, f in enumerate(files, 1):
            try:
                n = import_file(db, f)
                total += n
                print(f"[{i}/{len(files)}] {os.path.basename(f)} → {n} 行")
            except Exception as e:
                db.rollback()
                print(f"[{i}/{len(files)}] {os.path.basename(f)} ✗ 失败：{e}")
    finally:
        db.close()
    print(f"固定资产导入完成：{len(files)} 文件，共 {total} 行（fixed_assets + devices + aliases）。")


if __name__ == "__main__":
    main()

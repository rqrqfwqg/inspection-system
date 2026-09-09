"""定向修复：仅对 import_fixed_assets 中 row_count=0 的批次回填「文件级行数」。

背景：固定资产按「移交编号」去重，重复文件（如 E11地块2号楼 的 (未修改)/0210/补充12/
地下室/室外 等同名变体）的所有行都撞上已导入编号，被合并进规范批次，故 source_batch_id
指向规范批次、变体批次关联计数为 0。用「关联计数」反查会把它们误清零。

本脚本改用「文件级行数」口径（与 import_fixed_assets.py 的 count 一致）：只读取源文件、
用相同跳过规则统计数据行数并写回 row_count，file_count 统一=1。**不修改 fixed_assets 任何数据**。

安全：仅处理当前 row_count=0 的批次；只读 xlsx；仅 UPDATE import_batches 一行两列。
用法：venv/Scripts/python.exe scripts/fix_zero_fixedasset_batches.py
"""
import os
import sys

BACKEND_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, BACKEND_DIR)

import openpyxl
from database import engine, ImportBatch
from sqlalchemy.orm import sessionmaker

Session = sessionmaker(bind=engine)
DEFAULT_DIR = r"D:/机场业务/T3GTC/03_资产清单/公管分公司资产清单/公管分公司资产清单/"


def to_str(v):
    if v is None:
        return ""
    if isinstance(v, float):
        return str(int(v)) if v == int(v) else str(v)
    return str(v).strip()


def detect_header(ws):
    for r in range(1, 9):
        cells = [ws.cell(row=r, column=c).value for c in range(1, 24)]
        for c in cells:
            if isinstance(c, str) and "移交编号" in c:
                return r, r + 1
    return 4, 5


def count_file_rows(path: str) -> int:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    _, data_start = detect_header(ws)
    n = 0
    for row in ws.iter_rows(min_row=data_start, max_col=23, values_only=True):
        if row is None or all(v is None for v in row):
            continue
        vals = list(row) + [None] * (23 - len(row))
        tn = to_str(vals[1])
        an = to_str(vals[6])
        bm = to_str(vals[8])
        if not tn or tn == "移交编号" or "名称" in tn:
            continue
        if "（" in tn or "）" in tn:
            continue
        if not an and not bm:
            continue
        n += 1
    wb.close()
    return n


def main():
    db = Session()
    try:
        zeros = db.query(ImportBatch).filter(
            ImportBatch.imported_by == "import_fixed_assets",
            (ImportBatch.row_count == 0) | (ImportBatch.row_count.is_(None)),
        ).all()
        print(f"待修复的 import_fixed_assets 零值批次：{len(zeros)} 条")
        fixed = 0
        for b in zeros:
            path = os.path.join(DEFAULT_DIR, b.file_name) if b.file_name else None
            if not path or not os.path.exists(path):
                print(f"  · #{b.id} {b.file_name}: 源文件缺失，跳过")
                continue
            try:
                rc = count_file_rows(path)
            except Exception as e:
                print(f"  · #{b.id} {b.file_name}: 读取失败 {e}")
                continue
            b.row_count = rc
            b.file_count = 1
            print(f"  · #{b.id} {b.file_name}: row_count={rc}")
            fixed += 1
        db.commit()
        print(f"修复完成 ✓（{fixed} 条已写回；fixed_assets 数据未改动）")
    finally:
        db.close()


if __name__ == "__main__":
    main()

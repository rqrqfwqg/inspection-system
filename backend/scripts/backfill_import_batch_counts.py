"""一次性回填 import_batches 的 file_count / row_count（P1 数据质量修复）。

幂等、安全：
  - file_count 统一 = 1（设计约定：每文件 = 1 条批次）。
  - row_count：
      * import_fixed_assets 且已非 0 → 不动（保留导入脚本写入的文件级行数）；
        若为 0（重复文件去重导致关联计数为 0）→ 按源文件扫描真实文件级行数回填，
        **不修改 fixed_assets 任何数据**。
      * import_ba + 段=问题清单 → BaProblem.source_batch_id 计数
      * import_ba + 其它      → DeviceArchive.source_batch_id 计数
      * seed_rooms          → Room 总数（机房仅单一批次）

注意：import_fixed_assets 的重复文件变体（如 E11地块2号楼 的 (未修改)/0210/补充12 等）
其行全部并入规范批次，关联计数为 0；此处用「文件级行数」口径回填，与 import_fixed_assets.py
的 count 一致，避免误清零。如需重新扫描这些文件，可直接跑 fix_zero_fixedasset_batches.py。

用法：
  venv/Scripts/python.exe scripts/backfill_import_batch_counts.py
"""
import os
import sys

BACKEND_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, BACKEND_DIR)

import openpyxl
from database import engine, ImportBatch, FixedAsset, BaProblem, DeviceArchive, Room
from sqlalchemy.orm import sessionmaker

Session = sessionmaker(bind=engine)
DEFAULT_DIR = r"D:/机场业务/T3GTC/03_资产清单/公管分公司资产清单/公管分公司资产清单/"


def to_str(v):
    if v is None:
        return ""
    if isinstance(v, float):
        return str(int(v)) if v == int(v) else str(v)
    return str(v).strip()


def detect_header(ws):
    for r in range(1, 9):
        cells = [ws.cell(row=r, column=c).value for c in range(1, 24)]
        for c in cells:
            if isinstance(c, str) and "移交编号" in c:
                return r, r + 1
    return 4, 5


def count_file_rows(path: str) -> int:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    _, data_start = detect_header(ws)
    n = 0
    for row in ws.iter_rows(min_row=data_start, max_col=23, values_only=True):
        if row is None or all(v is None for v in row):
            continue
        vals = list(row) + [None] * (23 - len(row))
        tn = to_str(vals[1]); an = to_str(vals[6]); bm = to_str(vals[8])
        if not tn or tn == "移交编号" or "名称" in tn:
            continue
        if "（" in tn or "）" in tn:
            continue
        if not an and not bm:
            continue
        n += 1
    wb.close()
    return n


def main():
    db = Session()
    try:
        rows = db.query(ImportBatch).all()
        print(f"待回填批次：{len(rows)} 条")
        for b in rows:
            b.file_count = 1
            if b.imported_by == "import_fixed_assets":
                if b.row_count and b.row_count > 0:
                    rc = b.row_count  # 已正确（导入脚本文件级行数），保留
                else:
                    path = os.path.join(DEFAULT_DIR, b.file_name) if b.file_name else None
                    if path and os.path.exists(path):
                        rc = count_file_rows(path)
                    else:
                        rc = 0
            elif b.imported_by == "seed_rooms":
                rc = db.query(Room).count()
            elif b.imported_by == "import_ba":
                if b.segment_name == "问题清单":
                    rc = db.query(BaProblem).filter(BaProblem.source_batch_id == b.id).count()
                else:
                    rc = db.query(DeviceArchive).filter(DeviceArchive.source_batch_id == b.id).count()
            else:
                rc = b.row_count or 0
            b.row_count = rc
            print(f"  · #{b.id} {b.file_name or b.segment_name}: file_count=1 row_count={rc}")
        db.commit()
        print("回填完成 ✓")
    finally:
        db.close()


if __name__ == "__main__":
    main()
